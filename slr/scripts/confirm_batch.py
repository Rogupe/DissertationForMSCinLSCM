"""Write FINAL decisions for a confirmed batch and refresh PRISMA_log counters.
usage: python3 confirm_batch.py <workbook> <proposals_module> <batch> <date> [overrides_module]
Overrides module (optional) defines O = {id: (decision, code, final_justification)} for DUDAs
resolved with Rodrigo, or for rows he changed at confirmation.
"""
import os as _os, sys as _sys; _sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
import openpyxl, sys, importlib, collections, datetime
from openpyxl.styles import Alignment, Font, PatternFill
path, modname, batch, date = sys.argv[1:5]
P = importlib.import_module(modname).P
O = importlib.import_module(sys.argv[5]).O if len(sys.argv) > 5 else {}
REASON = {'E1': 'No digital twin/model', 'E2': 'No supply chain scope', 'E3': 'No AI/ML component',
          'E4': 'Wrong publication type', 'E5': 'Wrong publication type', 'E6': 'Duplicate'}
E5_REASON = {'idioma': 'Not in English', 'gris': 'Not peer-reviewed'}
wb = openpyxl.load_workbook(path)
ws = wb['Screening']
from screening_cols import cols
C = cols(ws)
rowof = {ws.cell(r, C('ID')).value: r for r in range(2, ws.max_row + 1)}
written = collections.Counter()
for rid, (dec, code, just, make) in P.items():
    r = rowof[rid]
    assert ws.cell(r, C('Lote')).value == int(batch), f'{rid} not in batch {batch}'
    if rid in O:
        ov = O[rid]
        dec, code, just = ov[0], ov[1], ov[2]
        if len(ov) > 3:
            make = ov[3]
            ws.cell(r, C('Nota Make')).value = make or None
    if code == 'DUDA':
        k, l = 'Unsure', None
    elif code == 'I':
        k, l = 'Include', None
    else:
        k = 'Exclude'
        l = REASON[code]
        if code == 'E5':
            for key, val in E5_REASON.items():
                if key in just.lower():
                    l = val
    note = (f'[Make] ' if make else '') + just
    ws.cell(r, C('Decision')).value = k
    ws.cell(r, C('Exclusion reason')).value = l
    ws.cell(r, C('Notes')).value = note
    ws.cell(r, C('Código final')).value = code
    for col in (C('Decision'), C('Exclusion reason'), C('Notes'), C('Código final')):
        ws.cell(r, col).alignment = Alignment(vertical='top')
    written[code] += 1

# ---- PRISMA_log: B10 literal (yellow input cell) + progress block below row 22 ----
pl = wb['PRISMA_log']
allK = [ws.cell(r, C('Decision')).value for r in range(2, ws.max_row + 1)]
allO = [ws.cell(r, C('Código final')).value for r in range(2, ws.max_row + 1)]
n_excl = sum(1 for v in allK if v == 'Exclude')
n_incl = sum(1 for v in allK if v == 'Include')
n_uns = sum(1 for v in allK if v == 'Unsure')
n_done = n_excl + n_incl + n_uns
pl['B10'].value = n_excl
codes = collections.Counter(v for v in allO if v)
batches = sorted({ws.cell(r, C('Lote')).value for r in range(2, ws.max_row + 1)
                  if ws.cell(r, C('Lote')).value and ws.cell(r, C('Decision')).value})
block = [
    ('TA SCREENING PROGRESS (running tally; B10 above is the confirmed-exclusion count so far)', None),
    (f'Updated', date),
    ('Records screened and confirmed', n_done),
    ('  of which Include (sought for full text)', n_incl),
    ('  of which Exclude', n_excl),
    ('  of which Unsure (DUDA pending)', n_uns),
    ('Records still to screen', 712 - n_done),
    ('Confirmed batches', ', '.join(str(b) for b in batches)),
    ('Code I  (include for full text)', codes.get('I', 0)),
    ('Code E1 (no digital twin / model)', codes.get('E1', 0)),
    ('Code E2 (no supply chain / logistics framing)', codes.get('E2', 0)),
    ('Code E3 (no AI / ML)', codes.get('E3', 0)),
    ('Code E4 (review; retained for backward snowballing)', codes.get('E4', 0)),
    ('Code E5 (out of range / not English / grey)', codes.get('E5', 0)),
    ('Code E6 (undetected duplicate)', codes.get('E6', 0)),
    ('Make-annotated includes', sum(1 for r in range(2, ws.max_row + 1)
                                    if ws.cell(r, C('Nota Make')).value == 'Make' and ws.cell(r, C('Decision')).value == 'Include')),
    ('E3 retained as background (Antecedente C8 pool)', sum(1 for r in range(2, ws.max_row + 1)
                                    if ws.cell(r, C('Código final')).value == 'E3' and 'Antecedente C8' in (ws.cell(r, C('Notes')).value or ''))),
]
start = 24
hdr_font = Font(bold=True, color='FFFFFF'); hdr_fill = PatternFill('solid', fgColor='1E2761')
for i, (a, b) in enumerate(block):
    ra = start + i
    pl.cell(ra, 1).value = a; pl.cell(ra, 2).value = b
    if i == 0:
        pl.cell(ra, 1).font = hdr_font; pl.cell(ra, 1).fill = hdr_fill
        pl.cell(ra, 2).fill = hdr_fill
import pools
pools.rebuild(wb)
wb.save(path)
print('batch', batch, 'written:', dict(written), '| totals: done', n_done, 'incl', n_incl, 'excl', n_excl, 'unsure', n_uns)
