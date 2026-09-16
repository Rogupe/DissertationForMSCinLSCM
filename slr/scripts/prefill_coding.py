"""Pre-populate the Coding sheet with the confirmed includes (ID, Harvard authors, year, venue, DOI, note hint).
usage: python3 prefill_coding.py <workbook>
"""
import os as _os, sys as _sys; _sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
import openpyxl, re, sys
from openpyxl.styles import Alignment
path = sys.argv[1]
wb = openpyxl.load_workbook(path)
ws = wb['Screening']; cd = wb['Coding']
from screening_cols import cols
C = cols(ws)

def harvard(authors):
    out = []
    for a in [x.strip() for x in re.split(r';', authors or '') if x.strip()]:
        if ',' in a:                      # WoS style "Lignell, A" / "Sánchez-Banderas, RM"
            sur, ini = [p.strip() for p in a.split(',', 1)]
            ini = re.sub(r'\.', '', ini)
            ini = ''.join(ch + '.' if ch.isalpha() else ch for ch in ini)
        else:                             # Scopus style "Guzmán E." / "Kadum Al-Addal S.I." / "Ten Hompel M."
            toks = a.split()
            if len(toks) > 1 and re.fullmatch(r'[A-Za-zÀ-ÿ\.\-]+\.', toks[-1]) and len(toks[-1].replace('.', '').replace('-', '')) <= 4:
                sur, ini = ' '.join(toks[:-1]), toks[-1]
            else:
                sur, ini = a, ''
        out.append(f'{sur}, {ini}'.strip().rstrip(','))
    if not out:
        return ''
    if len(out) == 1:
        return out[0]
    return ', '.join(out[:-1]) + ' and ' + out[-1]

rows = []
for r in range(2, ws.max_row + 1):
    if ws.cell(r, C('Decision')).value != 'Include':
        continue
    rid, auth, year, venue, doi, note, make = (ws.cell(r, C('ID')).value, ws.cell(r, C('Authors')).value, ws.cell(r, C('Year')).value,
                                              ws.cell(r, C('Venue')).value, ws.cell(r, C('DOI')).value, ws.cell(r, C('Notes')).value or '', ws.cell(r, C('Nota Make')).value)
    hint = []
    if make == 'Make':
        hint.append('TA: Make-annotated production/shop-floor twin with supply chain framing')
    m = re.search(r'Auditoría final 2026-09-06: se mantiene I; (.*)$', note)
    if m:
        hint.append('TA flag: ' + m.group(1).split(' | ')[0][:300])
    elif 'sensibilidad' in note.lower() or 'verificar' in note.lower() or 'confirmar' in note.lower():
        hint.append('TA: included under the sensitivity rule; check the flagged point in the Screening note')
    rows.append((rid, harvard(auth), year, venue, doi) + (None,) * 17 + ('; '.join(hint) or None,))
rows.sort()
# wipe rows 2..max and write
for r in range(2, cd.max_row + 1):
    for c in range(1, cd.max_column + 1):
        cd.cell(r, c).value = None
for i, row in enumerate(rows, start=2):
    for c, v in enumerate(row, start=1):
        cell = cd.cell(i, c, value=v)
        cell.alignment = Alignment(vertical='top', wrap_text=c in (2, 4, 19, 20, 23))
cd.auto_filter.ref = f'A1:W{len(rows) + 1}'
cd.freeze_panes = 'B2'
wb.save(path)
print('Coding rows written:', len(rows), '| sample:', rows[0][:5], '|', rows[-1][:5])
