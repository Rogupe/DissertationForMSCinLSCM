"""Create the Fulltext_eligibility sheet (retrieval + full-text decision log) for the confirmed includes.
usage: python3 build_fulltext_sheet.py <workbook>
"""
import os as _os, sys as _sys; _sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
import openpyxl, sys
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
path = sys.argv[1]
wb = openpyxl.load_workbook(path)
ws = wb['Screening']
from screening_cols import cols
C = cols(ws)
name = 'Fulltext_eligibility'
if name in wb.sheetnames:
    del wb[name]
ft = wb.create_sheet(name)
hdr = ['ID', 'Authors', 'Year', 'Title', 'Venue', 'Doc type', 'DOI', 'DOI link', 'Make (TA)', 'TA flag',
       'Retrieved (Y/N)', 'PDF file name', 'FT decision', 'FT exclusion reason', 'FT note (quote page/section)', 'FT date']
ft.append(hdr)
rows = []
for r in range(2, ws.max_row + 1):
    if ws.cell(r, C('Decision')).value != 'Include':
        continue
    note = ws.cell(r, C('Notes')).value or ''
    flag = ''
    if 'Auditoría final 2026-09-06: se mantiene I; ' in note:
        flag = note.split('Auditoría final 2026-09-06: se mantiene I; ', 1)[1].split(' | ')[0]
    elif any(k in note.lower() for k in ('sensibilidad', 'verificar', 'confirmar')):
        flag = 'sensibilidad: ver nota de cribado'
    doi = (ws.cell(r, C('DOI')).value or '').strip()
    rows.append([ws.cell(r, C('ID')).value, ws.cell(r, C('Authors')).value, ws.cell(r, C('Year')).value, ws.cell(r, C('Title')).value, ws.cell(r, C('Venue')).value,
                 ws.cell(r, C('Doc type')).value, doi, f'https://doi.org/{doi}' if doi else '', 'Make' if ws.cell(r, C('Nota Make')).value == 'Make' else '',
                 flag, None, None, None, None, None, None])
rows.sort(key=lambda x: x[0])
for row in rows:
    ft.append(row)
n = len(rows) + 1
for c in range(1, len(hdr) + 1):
    cell = ft.cell(1, c); cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1E2761')
    cell.alignment = Alignment(vertical='top', wrap_text=True)
for row in ft.iter_rows(min_row=2, max_row=n):
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=True)
for c in range(11, 17):  # yellow input cells
    for r in range(2, n + 1):
        ft.cell(r, c).fill = PatternFill('solid', fgColor='FFF9C4')
widths = [8, 28, 6, 48, 26, 12, 26, 30, 8, 40, 10, 22, 12, 30, 40, 11]
for i, w in enumerate(widths, start=1):
    ft.column_dimensions[get_column_letter(i)].width = w
dv1 = DataValidation(type='list', formula1='"Y,N"', allow_blank=True); dv1.add(f'K2:K{n}')
dv2 = DataValidation(type='list', formula1='"Include,Exclude"', allow_blank=True); dv2.add(f'M2:M{n}')
dv3 = DataValidation(type='list', formula1='"No AI/ML component,No digital twin / digital model,No supply chain / logistics scope,Not peer-reviewed / low-quality venue,Other (note reason)"', allow_blank=True); dv3.add(f'N2:N{n}')
for dv in (dv1, dv2, dv3):
    ft.add_data_validation(dv)
ft.freeze_panes = 'B2'
ft.auto_filter.ref = f'A1:P{n}'
# summary block feeding PRISMA_log rows 12-19
s = n + 2
ft.cell(s, 1).value = 'SUMMARY (copy into PRISMA_log when the full-text stage closes)'; ft.cell(s, 1).font = Font(bold=True)
lines = [
    ('Reports sought for retrieval', f'=COUNTA(A2:A{n})'),
    ('Reports not retrievable (Retrieved = N)', f'=COUNTIF(K2:K{n},"N")'),
    ('Reports assessed (Retrieved = Y)', f'=COUNTIF(K2:K{n},"Y")'),
    ('Excluded: no AI/ML component', f'=COUNTIF(N2:N{n},"No AI/ML component")'),
    ('Excluded: no digital twin / digital model', f'=COUNTIF(N2:N{n},"No digital twin / digital model")'),
    ('Excluded: no supply chain / logistics scope', f'=COUNTIF(N2:N{n},"No supply chain / logistics scope")'),
    ('Excluded: not peer-reviewed / low-quality venue', f'=COUNTIF(N2:N{n},"Not peer-reviewed / low-quality venue")'),
    ('Excluded: other', f'=COUNTIF(N2:N{n},"Other (note reason)")'),
    ('Included after full text (FT decision = Include)', f'=COUNTIF(M2:M{n},"Include")'),
    ('Pending (retrieved, no FT decision yet)', f'=COUNTIF(K2:K{n},"Y")-COUNTIF(M2:M{n},"Include")-COUNTIF(M2:M{n},"Exclude")'),
]
for i, (a, b) in enumerate(lines, start=1):
    ft.cell(s + i, 1).value = a; ft.cell(s + i, 2).value = b
wb.save(path)
print('Fulltext_eligibility rows:', len(rows), '| flagged:', sum(1 for r in rows if r[9]), '| Make:', sum(1 for r in rows if r[8]))
