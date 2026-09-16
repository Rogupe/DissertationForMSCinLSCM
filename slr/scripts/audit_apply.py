"""Apply the final-audit adjudication (audit_recodes.py) to the screening workbook.
usage: python3 audit_apply.py <workbook>  (writes in place; keeps a .bak copy)
"""
# Repository copy (2026-09-16): quote verification needs the working workbook WITH the Abstract column and keywords_by_id.json (both derived from the raw Scopus/WoS exports, not distributed); on the repository copy of the workbook the quotes cannot be re-verified.

import os as _os, sys as _sys; _sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
import openpyxl, json, re, sys, shutil, unicodedata, collections, os
from copy import copy
from openpyxl.styles import Alignment, Font, PatternFill
from audit_recodes import R, V, POOLV, LINKS, D
import pools

path = sys.argv[1]
shutil.copy(path, path + '.bak')
kw = json.load(open('keywords_by_id.json')) if os.path.exists('keywords_by_id.json') else {}  # keyword file derived from the raw exports (not in the repository)
wb = openpyxl.load_workbook(path)
ws = wb['Screening']
from screening_cols import cols
C = cols(ws)
rowof = {ws.cell(r, C('ID')).value: r for r in range(2, ws.max_row + 1)}
REASON = {'E1': 'No digital twin/model', 'E2': 'No supply chain scope', 'E3': 'No AI/ML component',
          'E4': 'Wrong publication type', 'E5': 'Wrong publication type', 'E6': 'Duplicate'}

def norm(s):
    s = s.replace('’', "'").replace('‘', "'").replace('“', '"').replace('”', '"')
    s = s.replace('\u2013', '-').replace('\u2014', '-').replace(' ', ' ')
    s = unicodedata.normalize('NFKC', s)
    return re.sub(r'\s+', ' ', s).casefold()

# ---- 1. verify every quote verbatim -----------------------------------------------------------
problems = 0
for rid, (code, value, make, just) in R.items():
    r = rowof[rid]
    hay = norm(' || '.join([ws.cell(r, C('Title')).value or '', (ws.cell(r, C('Abstract')).value or '') if C.has('Abstract') else '']
                           + [v for k, v in kw.get(rid, {}).items() if isinstance(v, str)]))
    for q in re.findall(r'"([^"]+)"', just):
        if norm(q) not in hay:
            print('QUOTE NOT FOUND', rid, '->', q); problems += 1
assert problems == 0, f'{problems} quote problems'
print('quotes verified for', len(R), 'records')

def strip_prefixes(note):
    """Remove [Make] / [Antecedente C8: x] prefixes from a note being demoted to 'original' text."""
    tags = []
    changed = True
    while changed:
        changed = False
        m = re.match(r'\[(Make|Antecedente C8: (?:alto|medio|bajo))\]\s*', note)
        if m:
            tags.append(m.group(1)); note = note[m.end():]; changed = True
    return note, tags

def set_align(r):
    for col in (C('Decision'), C('Exclusion reason'), C('Notes'), C('Código final'), C('Nota Make')):
        ws.cell(r, col).alignment = Alignment(vertical='top')

log = collections.Counter(); changes = []
# ---- 2. recodes and flags ---------------------------------------------------------------------
for rid, (code, value, make, just) in R.items():
    r = rowof[rid]
    old_code = ws.cell(r, C('Código final')).value
    old_note = ws.cell(r, C('Notes')).value or ''
    old_make = ws.cell(r, C('Nota Make')).value
    if code == 'I':
        assert old_code == 'I', (rid, old_code)
        note = old_note
        if make == 'Make' and old_make != 'Make':
            ws.cell(r, C('Nota Make')).value = 'Make'
            note = '[Make] ' + note
        note = note + f' | Auditoría final {D}: se mantiene I; ' + just
        ws.cell(r, C('Notes')).value = note
        log['kept-flagged'] += 1
        changes.append((rid, old_code, 'I', 'flag' + (' +Make' if make == 'Make' else '')))
        set_align(r); continue
    # exclusion recode
    body, tags = strip_prefixes(old_note)
    body = body.replace('Antecedente C8: medio', 'pool de antecedentes (valor medio)')
    tag_txt = ''
    if 'Make' in tags:
        tag_txt += ' (anotado Make en el cribado inicial, retirado)'
    for t in tags:
        if t.startswith('Antecedente'):
            tag_txt += f' (antes en el pool de antecedentes, valor {t.split(": ")[1]})'
    if code == 'E6':
        primary = value
        note = f'Recodificado {D} en la auditoría final ({old_code}→E6, informe compañero de {primary}): {just} | Nota original ({old_code}){tag_txt}: {body}'
    elif code == 'E4':
        note = f'Recodificado {D} en la auditoría final ({old_code}→E4): {just} Valor para snowballing: {value}. | Nota original ({old_code}){tag_txt}: {body}'
    elif code == 'E3':
        pre = f'[Antecedente C8: {value}] ' if value else ''
        note = f'{pre}Recodificado {D} en la auditoría final ({old_code}→E3): {just} | Nota original ({old_code}){tag_txt}: {body}'
    else:
        note = f'Recodificado {D} en la auditoría final ({old_code}→{code}): {just} | Nota original ({old_code}){tag_txt}: {body}'
    ws.cell(r, C('Decision')).value = 'Exclude'
    ws.cell(r, C('Exclusion reason')).value = REASON[code]
    ws.cell(r, C('Notes')).value = note
    ws.cell(r, C('Código final')).value = code
    if make == '':
        ws.cell(r, C('Nota Make')).value = None
    elif old_make == 'Make':
        ws.cell(r, C('Nota Make')).value = None  # an excluded record carries no Make annotation
    log[f'{old_code}->{code}'] += 1
    changes.append((rid, old_code, code, value))
    set_align(r)

# ---- 3. snowballing value changes ------------------------------------------------------------
for rid, (old, new, why) in V.items():
    r = rowof[rid]
    assert ws.cell(r, C('Código final')).value == 'E4', rid
    note = ws.cell(r, C('Notes')).value
    key = f'Valor para snowballing: {old}'
    assert note.count(key) == 1, (rid, note.count(key))
    ws.cell(r, C('Notes')).value = note.replace(key, f'Valor para snowballing: {new} (ajustado en la auditoría final {D}, antes {old}: {why})')
    log['E4 value'] += 1
    changes.append((rid, f'E4 {old}', f'E4 {new}', why))
for rid, (old, new, why) in POOLV.items():
    r = rowof[rid]
    note = ws.cell(r, C('Notes')).value
    key = f'[Antecedente C8: {old}]'
    assert note.startswith(key), rid
    ws.cell(r, C('Notes')).value = f'[Antecedente C8: {new}]' + note[len(key):] + f' | Auditoría final {D}: valor del pool ajustado de {old} a {new} ({why}).'
    log['pool value'] += 1
    changes.append((rid, f'E3 pool {old}', f'E3 pool {new}', why))

# ---- 4. cross-links -------------------------------------------------------------------------
for rid, txt in LINKS.items():
    r = rowof[rid]
    ws.cell(r, C('Notes')).value = (ws.cell(r, C('Notes')).value or '') + f' | Auditoría final {D}: {txt}'
    log['link'] += 1

# ---- 5. reword a pool mention inside an include's note --------------------------------------
r = rowof['R0664']
n = ws.cell(r, C('Notes')).value
assert 'Antecedente C8: medio' in n
ws.cell(r, C('Notes')).value = n.replace('alternativa E3 con Antecedente C8: medio', 'alternativa E3 en el pool de antecedentes (valor medio)')

# ---- 6. PRISMA_log ---------------------------------------------------------------------------
pl = wb['PRISMA_log']
allK = [ws.cell(r, C('Decision')).value for r in range(2, ws.max_row + 1)]
allO = [ws.cell(r, C('Código final')).value for r in range(2, ws.max_row + 1)]
n_excl = sum(1 for v in allK if v == 'Exclude'); n_incl = sum(1 for v in allK if v == 'Include')
n_uns = sum(1 for v in allK if v == 'Unsure'); n_done = n_excl + n_incl + n_uns
assert n_done == 712 and n_uns == 0
pl['B10'].value = n_excl
codes = collections.Counter(v for v in allO if v)
n_make = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, C('Nota Make')).value == 'Make' and ws.cell(r, C('Decision')).value == 'Include')
n_pool = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, C('Código final')).value == 'E3' and 'Antecedente C8' in (ws.cell(r, C('Notes')).value or ''))
block = [
    ('TA SCREENING PROGRESS (final; B10 above is the confirmed-exclusion count)', None),
    ('Updated', f'{D} (final audit applied)'),
    ('Records screened and confirmed', n_done),
    ('  of which Include (sought for full text)', n_incl),
    ('  of which Exclude', n_excl),
    ('  of which Unsure (DUDA pending)', n_uns),
    ('Records still to screen', 712 - n_done),
    ('Confirmed batches', '1 to 14, then final audit'),
    ('Code I  (include for full text)', codes.get('I', 0)),
    ('Code E1 (no digital twin / model)', codes.get('E1', 0)),
    ('Code E2 (no supply chain / logistics framing)', codes.get('E2', 0)),
    ('Code E3 (no AI / ML)', codes.get('E3', 0)),
    ('Code E4 (review; retained for backward snowballing)', codes.get('E4', 0)),
    ('Code E5 (out of range / not English / grey)', codes.get('E5', 0)),
    ('Code E6 (duplicate / companion report of an included study)', codes.get('E6', 0)),
    ('Make-annotated includes', n_make),
    ('E3 retained as background (Antecedente C8 pool)', n_pool),
]
hdr_font = Font(bold=True, color='FFFFFF'); hdr_fill = PatternFill('solid', fgColor='1E2761')
for i, (a, b) in enumerate(block):
    ra = 24 + i
    pl.cell(ra, 1).value = a; pl.cell(ra, 2).value = b
    if i == 0:
        pl.cell(ra, 1).font = hdr_font; pl.cell(ra, 1).fill = hdr_fill; pl.cell(ra, 2).fill = hdr_fill

# ---- 7. pools ---------------------------------------------------------------------------------
n_e4, n_c8 = pools.rebuild(wb)
wb.save(path)
print('changes:', dict(log))
print('totals: I', n_incl, 'excl', n_excl, dict(sorted(codes.items())), 'Make', n_make, 'pool', n_pool, 'E4 sheet', n_e4, 'C8 sheet', n_c8)
json.dump(changes, open('audit_changes.json', 'w'), ensure_ascii=False, indent=1)
