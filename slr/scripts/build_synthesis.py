"""Build the Synthesis sheet (descriptive tables + typology cross-tabs) and Included_studies from Coding.
usage: python3 build_synthesis.py <workbook>
"""
# Repository copy (2026-09-16): coding records read from slr/coding/ftcodes (override with SLR_FTCODES).

import os as _os, sys as _sys; _sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
import openpyxl, sys, collections, json, os
FTCODES = os.environ.get('SLR_FTCODES', os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'coding', 'ftcodes'))
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
path = sys.argv[1]
wb = openpyxl.load_workbook(path)
cd = wb['Coding']; ws = wb['Screening']
from screening_cols import cols
C = cols(ws)
hdr = [cd.cell(1, c).value for c in range(1, cd.max_column + 1)]
col = {h: i + 1 for i, h in enumerate(hdr)}
recs = {}
for f in os.listdir(FTCODES):
    d = json.load(open(os.path.join(FTCODES, f), encoding='utf-8')); recs[d['id']] = d
scr = {ws.cell(r, C('ID')).value: (ws.cell(r, C('Authors')).value, ws.cell(r, C('Year')).value, ws.cell(r, C('Title')).value, ws.cell(r, C('Venue')).value, ws.cell(r, C('DOI')).value, ws.cell(r, C('Nota Make')).value)
       for r in range(2, ws.max_row + 1)}
rows = []
for r in range(2, cd.max_row + 1):
    rid = cd.cell(r, 1).value
    if not rid or not cd.cell(r, col['Evidence type']).value:
        continue
    g = lambda h: cd.cell(r, col[h]).value
    rows.append(dict(id=rid, ev=g('Evidence type'), sector=g('Sector'), plan=g('SCOR: Plan'), source=g('SCOR: Source'), make=g('SCOR: Make'),
                     deliver=g('SCOR: Deliver'), ret=g('SCOR: Return'), e2e=g('End-to-end'), tech=g('AI technique'), role=g('AI role in twin'),
                     ext=g('Integration extent (1-5)'), hrole=g('HITL role'), hfun=g('HITL function'), typ=g('Use in typology (Y/N)'),
                     key=g('Key finding (own words)'), cond=g('Conditions under which HITL changes')))
prim = lambda v: (v or '').split(';')[0].strip()

def table(title, counter, total, start_row, sheet, note=None):
    sheet.cell(start_row, 1).value = title; sheet.cell(start_row, 1).font = Font(bold=True)
    sheet.cell(start_row + 1, 1).value = 'Category'; sheet.cell(start_row + 1, 2).value = 'n'; sheet.cell(start_row + 1, 3).value = '% of ' + str(total)
    for c in (1, 2, 3):
        sheet.cell(start_row + 1, c).font = Font(bold=True, color='FFFFFF'); sheet.cell(start_row + 1, c).fill = PatternFill('solid', fgColor='1E2761')
    i = start_row + 2
    for k, v in sorted(counter.items(), key=lambda x: (-x[1], str(x[0]))):
        sheet.cell(i, 1).value = k; sheet.cell(i, 2).value = v; sheet.cell(i, 3).value = round(100 * v / total, 1); i += 1
    if note:
        sheet.cell(i, 1).value = note; sheet.cell(i, 1).font = Font(italic=True); i += 1
    return i + 1

def crosstab(title, rowkey, colkey, data, start_row, sheet):
    rvals = sorted({rowkey(d) for d in data}); cvals = sorted({colkey(d) for d in data})
    sheet.cell(start_row, 1).value = title; sheet.cell(start_row, 1).font = Font(bold=True)
    for j, cv in enumerate(cvals, start=2):
        sheet.cell(start_row + 1, j).value = cv; sheet.cell(start_row + 1, j).font = Font(bold=True)
    sheet.cell(start_row + 1, len(cvals) + 2).value = 'Total'; sheet.cell(start_row + 1, len(cvals) + 2).font = Font(bold=True)
    for i, rv in enumerate(rvals, start=start_row + 2):
        sheet.cell(i, 1).value = rv; sheet.cell(i, 1).font = Font(bold=True)
        tot = 0
        for j, cv in enumerate(cvals, start=2):
            n = sum(1 for d in data if rowkey(d) == rv and colkey(d) == cv); sheet.cell(i, j).value = n or None; tot += n
        sheet.cell(i, len(cvals) + 2).value = tot
    i = start_row + 2 + len(rvals)
    sheet.cell(i, 1).value = 'Total'; sheet.cell(i, 1).font = Font(bold=True)
    for j, cv in enumerate(cvals, start=2):
        sheet.cell(i, j).value = sum(1 for d in data if colkey(d) == cv)
    sheet.cell(i, len(cvals) + 2).value = len(data)
    return i + 2

for name in ('Synthesis', 'Included_studies'):
    if name in wb.sheetnames:
        del wb[name]
sy = wb.create_sheet('Synthesis')
sy.cell(1, 1).value = f'SYNTHESIS OF THE {len(rows)} INCLUDED STUDIES (full-text stage closed 2026-09-06). Multi-valued fields are counted on their primary value.'
sy.cell(1, 1).font = Font(bold=True)
N = len(rows); r = 3
r = table('A. Evidence type', collections.Counter(d['ev'] for d in rows), N, r, sy)
r = table('B. Sector', collections.Counter((d['sector'] or '').split('(')[0].split('/')[0].strip() for d in rows), N, r, sy, 'Sector labels collapsed to their first term; see Coding for the full label.')
scor = collections.Counter()
for d in rows:
    for k, lab in (('plan', 'Plan'), ('source', 'Source'), ('make', 'Make'), ('deliver', 'Deliver'), ('ret', 'Return'), ('e2e', 'End-to-end')):
        if d[k] == 'Y': scor[lab] += 1
r = table('C. SCOR processes covered by the twin (studies may cover several)', scor, N, r, sy)
r = table('D. AI technique (primary)', collections.Counter(prim(d['tech']) for d in rows), N, r, sy)
r = table('E. AI role in the twin (primary)', collections.Counter(prim(d['role']) for d in rows), N, r, sy)
r = table('F. Integration extent', collections.Counter(d['ext'] for d in rows), N, r, sy)
r = table('G. HITL role', collections.Counter(d['hrole'] for d in rows), N, r, sy)
r = table('H. HITL function (primary)', collections.Counter(prim(d['hfun']) for d in rows), N, r, sy)
r = table('I. Usable in the typology', collections.Counter(d['typ'] for d in rows), N, r, sy)
r = crosstab('J. Integration extent x HITL role (all included studies)', lambda d: d['ext'], lambda d: d['hrole'], rows, r, sy)
typ = [d for d in rows if d['typ'] == 'Y']
r = crosstab(f'K. Integration extent x HITL role (typology subset, n={len(typ)})', lambda d: d['ext'], lambda d: d['hrole'], typ, r, sy)
r = crosstab('L. AI role (primary) x HITL function (primary), typology subset', lambda d: prim(d['role']), lambda d: prim(d['hfun']), typ, r, sy)
r = crosstab('M. Evidence type x Integration extent', lambda d: d['ev'], lambda d: d['ext'], rows, r, sy)
# ID lists per cell of table K for traceability
sy.cell(r, 1).value = 'N. Study IDs per cell of table K (integration extent | HITL role)'; sy.cell(r, 1).font = Font(bold=True); r += 1
cells = collections.defaultdict(list)
for d in typ:
    cells[(d['ext'], d['hrole'])].append(d['id'])
for (e, h), ids in sorted(cells.items()):
    sy.cell(r, 1).value = e; sy.cell(r, 2).value = h; sy.cell(r, 3).value = ', '.join(sorted(ids)); r += 1
# ---- O. reports not retrieved, by source (for the PRISMA diagram) ----
PUB = {'10.1109': 'IEEE Xplore conference proceedings (not subscribed by Hull or Tec; no open copy found)',
       '10.1007': 'Springer book chapters (IFIP AICT, LNNS, LNEE, LNME, Proc. Math. & Stat.; not subscribed; R0233 and R0253 embargoed on HAL until 2027-01-01)',
       '10.1108': 'Emerald (Kybernetes)', '10.1117': 'SPIE Proceedings', '10.1061': 'ASCE Library (CRC 2022)'}
nr = collections.defaultdict(list)
for d in recs.values():
    if d['status'] == 'NOT_RETRIEVED':
        nr[PUB.get(d['doi'].split('/')[0], d['doi'].split('/')[0])].append(d['id'])
sy.cell(r, 1).value = f'O. Reports not retrieved (n = {sum(len(v) for v in nr.values())}), by source, for the PRISMA diagram'; sy.cell(r, 1).font = Font(bold=True); r += 1
for lab, ids in sorted(nr.items(), key=lambda x: -len(x[1])):
    sy.cell(r, 1).value = lab; sy.cell(r, 2).value = len(ids); sy.cell(r, 3).value = ', '.join(sorted(ids)); r += 1
sy.cell(r, 1).value = 'Total'; sy.cell(r, 2).value = sum(len(v) for v in nr.values()); r += 2
# ---- P. included studies outside the typology subset, with the adjudicated reason ----
NOTYP = {'R0054': 'Human role not described; the SimPy "Level-1 twin" only stress-tests ML+OR policies offline',
         'R0057': 'No human appears anywhere in the closed GA-MARL loop (role and function "Not discussed")',
         'R0094': 'Conceptual proposal with nothing implemented; human role only implicit (practitioners use the surrogate)',
         'R0118': 'Twin only asserted (no description, data or validation); human role ambiguous ("operators or autonomous systems")',
         'R0119': 'Human role not described; twin predicts shelf life, decisions left to unspecified users',
         'R0173': 'Conceptual five-layer architecture; AI named only by families, human only as AR/VR interface',
         'R0174': 'Internally inconsistent results and reference list (quality doubts); no human in the loop',
         'R0249': 'Human role only inferred (operators currently decide; RL "without explicit human expertise"); RL algorithm unspecified',
         'R0629': 'AI limited to data mining and fuzzy rules without a named method; human role never discussed'}
out = [d for d in rows if d['typ'] != 'Y']
sy.cell(r, 1).value = f'P. Included studies outside the typology subset (use_in_typology = N; n = {len(out)}) and reason'; sy.cell(r, 1).font = Font(bold=True); r += 1
for d in sorted(out, key=lambda x: x['id']):
    sy.cell(r, 1).value = d['id']; sy.cell(r, 2).value = f"{d['ext']} | {d['hrole']} | {d['hfun']}"; sy.cell(r, 3).value = NOTYP.get(d['id'], 'see Coding notes'); r += 1
sy.cell(r, 1).value = ('Criterion: a study enters the typology subset only when the full text gives a legible AI-role + human-role configuration '
                       '(who decides, approves, sets or interprets); studies whose human role is undescribed, whose twin or AI is only asserted, '
                       'or whose reported results are internally inconsistent are kept as included studies (they count in the descriptive tables A-I) '
                       'but not in the typology cross-tabs (tables K, L, N).')
sy.cell(r, 1).font = Font(italic=True); r += 2
# ---- Q. reconciliation of table K ----
ext = collections.Counter(d['ext'] for d in typ); fun = collections.Counter(prim(d['hfun']) for d in typ)
sec = collections.Counter(x.strip() for d in typ for x in (d['hfun'] or '').split(';')[1:] if x.strip())
parts = []
for e in sorted(ext):
    sub = collections.Counter(d['hrole'] for d in typ if d['ext'] == e)
    parts.append(f"level {e[0]} = {ext[e]} ({', '.join(f'{v} {k.lower()}' for k, v in sub.most_common())})")
sy.cell(r, 1).value = (f'Q. Reconciliation of table K (n = {len(typ)}): ' + '; '.join(parts) + '. Primary human function (n = ' + str(len(typ)) + '): '
                       + ', '.join(f'{k.lower()} {v}' for k, v in fun.most_common()) + '; secondary functions: '
                       + (', '.join(f'{k.lower()} {v}' for k, v in sec.most_common()) or 'none') + '.')
r += 1
for c, w in zip('ABCDEFGH', [46, 24, 24, 24, 24, 24, 24, 12]):
    sy.column_dimensions[c].width = w
for row in sy.iter_rows():
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=True)
# Included_studies sheet
inc = wb.create_sheet('Included_studies')
inc.append(['ID', 'Authors', 'Year', 'Title', 'Venue', 'DOI', 'Make (TA)', 'Evidence type', 'Sector', 'SCOR (Y)', 'End-to-end', 'AI technique', 'AI role', 'Integration extent', 'HITL role', 'HITL function', 'Use in typology', 'Key finding'])
for d in sorted(rows, key=lambda x: x['id']):
    a, y, t, v, doi, make = scr[d['id']]
    scorlist = '/'.join(lab for k, lab in (('plan', 'Plan'), ('source', 'Source'), ('make', 'Make'), ('deliver', 'Deliver'), ('ret', 'Return')) if d[k] == 'Y')
    inc.append([d['id'], a, y, t, v, doi, make or '', d['ev'], d['sector'], scorlist, d['e2e'], d['tech'], d['role'], d['ext'], d['hrole'], d['hfun'], d['typ'], d['key']])
for c in range(1, 19):
    cell = inc.cell(1, c); cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1E2761')
for i, w in enumerate([8, 28, 6, 50, 26, 26, 8, 16, 22, 22, 8, 18, 22, 20, 14, 26, 8, 60], start=1):
    inc.column_dimensions[get_column_letter(i)].width = w
for row in inc.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=True)
inc.freeze_panes = 'B2'; inc.auto_filter.ref = inc.dimensions
wb.save(path)
print('included', N, '| typology subset', len(typ))
print('extent x hitl (typology):', {k: len(v) for k, v in sorted(cells.items())})
